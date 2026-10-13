"""NODO writes real PDF / Word / Excel / PowerPoint files from Markdown, and can read them back."""
from __future__ import annotations

import io
import zipfile

import pytest

from nodo.services import documents
from nodo.services.files import FileStore
from tests.test_brain_extra import box  # noqa: F401  (fixture)

MD = """# Preventivo sito

Ciao **Gaetano**, ecco il *preventivo* con accenti: è così, perché sì. €

## Voci

- Home page
- Contatti
1. Primo
2. Secondo

| Voce | Prezzo |
|------|--------|
| Home | 1.200,50 |
| Blog | 800 |

---

```
codice = 1
```

## Chiusura

Grazie.
"""


def test_pdf_has_a_pdf_signature_and_pages():
    data = documents.build("pdf", MD, "Preventivo")
    assert data.startswith(b"%PDF") and len(data) > 1500


def test_docx_and_pptx_are_real_office_zips():
    for fmt, marker in (("docx", "word/document.xml"), ("pptx", "ppt/presentation.xml")):
        data = documents.build(fmt, MD, "Preventivo")
        assert marker in zipfile.ZipFile(io.BytesIO(data)).namelist()


def test_xlsx_has_one_sheet_per_table_with_numbers():
    import openpyxl
    data = documents.build("xlsx", MD, "Preventivo")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.worksheets[0]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    flat = [v for r in rows for v in r]
    assert "Voce" in flat and 1200.5 in flat and 800 in flat


def test_pptx_has_a_slide_per_heading():
    from pptx import Presentation
    prs = Presentation(io.BytesIO(documents.build("pptx", MD, "Preventivo")))
    assert len(prs.slides) >= 3


def test_unknown_format_is_refused():
    with pytest.raises(ValueError):
        documents.build("exe", MD)


async def test_tool_creates_a_downloadable_document_and_reads_it_back(box, db, tmp_path):  # noqa: F811
    box.c.files = FileStore(tmp_path / "files")
    out = await box.call("create_document", {"name": "preventivo", "format": "docx", "content": MD,
                                             "title": "Preventivo sito"})
    assert out["created"] == "Preventivo sito" and out["url"].endswith("preventivo.docx")
    assert (tmp_path / "files" / "preventivo.docx").read_bytes()[:2] == b"PK"
    back = await box.call("read_file", {"name": "Preventivo sito"})
    assert "Gaetano" in back["content"] and "1.200,50" in back["content"]
    pdf = await box.call("create_document", {"name": "preventivo.pdf", "format": "pdf", "content": MD})
    assert pdf["url"].endswith("preventivo.pdf")
    assert (tmp_path / "files" / "preventivo.pdf").read_bytes().startswith(b"%PDF")
    assert "error" in await box.call("create_document", {"name": "x", "format": "exe", "content": "ciao"})
    assert "error" in await box.call("create_document", {"name": "x", "format": "pdf", "content": "  "})
